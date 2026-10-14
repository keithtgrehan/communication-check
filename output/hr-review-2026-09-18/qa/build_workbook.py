from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pathlib import Path

out=Path('/Users/keith/Documents/ChatGPT/communication-check/output/hr-review-2026-09-18')
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=Inches(.65); s.bottom_margin=Inches(.65)
s.left_margin=Inches(.75); s.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']; normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',22),('Heading 1',17),('Heading 2',12)]:
    st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True
    st.paragraph_format.space_before=Pt(10); st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.keep_with_next=True
doc.styles['Title'].paragraph_format.space_before=Pt(0)

def p(text='',bold=False):
    x=doc.add_paragraph(); x.add_run(text).bold=bold; return x

def h(text): doc.add_heading(text,2)
def field(question,answer='[Type here]',space=5):
    x=p(question,True); x.paragraph_format.keep_with_next=True
    x=p(answer); x.paragraph_format.space_after=Pt(space)
    x.runs[0].font.color.rgb=RGBColor.from_string('555555')

def link(text,url):
    x=doc.add_paragraph(); a=OxmlElement('w:hyperlink'); a.set(qn('r:id'),x.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); props=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'145A68'); props.append(c)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); props.append(u); r.append(props)
    t=OxmlElement('w:t'); t.text=text; r.append(t); a.append(r); x._p.append(a)

def page(title,time):
    doc.add_page_break(); doc.add_heading(title,1); p(time)

# Page 1
p('Communication Check review',True).style='Title'
p('A 45 minute workbook for recruitment and HR feedback',True)
p('Thanks for helping,\nKeith would like to learn where the tool helps, where it gets in the way, and what feels unclear. Short, honest answers are enough. You can skip a question or stop at any time.')
p('Open this file in Word or Google Docs. Replace [Type here] with your answers and keep your chosen options. Save your answers in this file as you go; the website does not collect this feedback. Return the completed file to Keith.')
p('Use invented messages and details only. Do not paste real candidate, employee or customer information into the tool or this workbook.')
link('Open Communication Check','https://communication-check.vercel.app/')
field('Review details','Date: [Type here]   Role or area of work: [Type here]\nDevice and browser: [Type here]\nCompleted: alone / with Keith observing / with help from someone else')
h('1 First impressions   3 minutes')
p('Open the website before reading further or getting a demonstration.')
field('What do you think the tool does and what would you click first?')
h('2 Try one sample   5 minutes')
p('The prototype is designed for customer onboarding follow-ups. Your recruitment and HR experience can help review its wording and ease of use. It checks some wording patterns; it cannot verify facts or agreements, and it can miss issues.')
p('Choose a built-in example and select Try a sample review. Look at both Clarity & commitments and Tone cues. A finding means a piece of wording the tool asks you to reconsider.')
field('Which sample did you try and was anything confusing?','Sample name: [Type here]\nWhat I tried / expected / saw instead: [Type here, or Nothing confusing]')

# Pages 2 and 3
for num,title,time,prompt in [
(3,'A message you would improve','14 minutes','Invent a short follow-up you would want to improve. It could follow an interview, explain a new starter’s next steps, or chase a response. Use 3 to 5 sentences that sound realistic to you.'),
(4,'A message you would already send','12 minutes','Invent a different short follow-up that you consider clear and appropriate. A firm, reasonable deadline is fine. Use 3 to 5 sentences; do not deliberately add a problem.')]:
    page(f'{num} {title}',time)
    p(prompt)
    h('Before using the checker')
    field('Who is writing to whom and what should happen next?','[Type here, including any relevant fictional context]')
    field('Paste your complete invented message here','[Type here before checking. Keep this original version unchanged in the workbook.]',8)
    field('What would you change before seeing the results and why?','[Type here, or Nothing — I would send it as written]')
    h('Now check the same message in both tabs')
    p('Paste it into the tool. Run Check draft under Clarity & commitments, then Check tone cues under Tone cues. Leave the tone question style on Collaborative. Record results before editing.')
    p('For each tab, record the finding that matters most to your judgment. Copy its quoted wording and review question. Choose Useful, Unnecessary or Unclear, and explain why. If there are no findings, write None. If you could not run it, write Not tested.')
    field('Clarity and commitments result','Quoted wording and review question: [Type here]\nMy judgment and reason: [Type here]')
    field('Tone cues result','Quoted wording and review question: [Type here]\nMy judgment and reason: [Type here]')
    field('Did it miss anything important or misunderstand the context?','[Quote the wording and explain, or write Nothing I noticed]')
    field('Did a finding change your mind about what to send?','[Yes / No / Unsure, and why. If yes, write the revised sentence.]')

# Page 4
page('5 Use the controls and give your verdict','11 minutes')
p('Use a finding from either message, or return to a built-in example if neither produced one. Try the actions below. If you get stuck for more than a minute, record where and move on.')
p('For each action write: Done alone / Done with help / Could not finish / Not tried. If you needed help or something felt wrong, add what happened.')
field('Which message or sample are you using?')
field('Keep one finding as written and then undo that decision','Result: [Type here]   Any help or difficulty: [Type here]')
field('Use Edit this wording and change one sentence yourself','Result: [Type here]   Any help or difficulty: [Type here]\nWhat happened to the earlier findings after the edit? [Type here]')
field('Run both checks again on your edited draft','Result: [Type here]   Any help or difficulty: [Type here]')
field('Use Copy draft and paste the result below','Result: [Type here]\nPasted text: [Type here]\nDid it include your latest edit? Yes / No / Not tested')
h('Your verdict')
field('When no wording is flagged what does that mean to you?')
field('What was the most useful moment and what should Keith fix first?','Most useful moment, or None: [Type here]\nOne thing to fix first, with an example: [Type here]')
field('Would you choose to use this again in your own work?','Yes / Maybe / No\nFor which situation, and what would you normally use or do instead? [Type here]')
field('Time spent and anything you skipped','Minutes: [Type here]   Skipped items and why: [Type here, or None]')
p('Save this file and send it back to Keith. Your answers are feedback for improving the prototype. They are not a public testimonial; Keith should ask separately before quoting you publicly.')

# Simple page numbers aid navigation in a four-page form.
f=s.footer.paragraphs[0]; f.alignment=2
r=f.add_run('Communication Check review  |  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('666666')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
for node in list(doc.styles.element.iter(qn('w:pBdr'))):
    node.getparent().remove(node)
for node in list(doc.element.iter(qn('w:pBdr'))):
    node.getparent().remove(node)
doc.core_properties.title='Communication Check review workbook'
doc.core_properties.subject='Recruitment and HR feedback in 45 minutes'
doc.core_properties.author='Keith Grehan'
doc.core_properties.keywords=''
doc.save(out/'Communication-Check-Review-Workbook.docx')
print(out/'Communication-Check-Review-Workbook.docx')
